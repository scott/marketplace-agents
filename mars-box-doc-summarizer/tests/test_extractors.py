from box_doc_summarizer.extractors import extract_bytes, is_supported
from helpers import FIXTURE


def _x(rel, **kw):
    p = FIXTURE / rel
    return extract_bytes(p.name, p.read_bytes(), **kw)


def test_markdown_text_csv():
    assert "guided setup" in _x("q3-roadmap.md").text.lower()
    assert "SSO" in _x("customer-feedback.txt").text
    assert "activation_rate" in _x("metrics.csv").text


def test_docx_and_pdf():
    assert "four engineers" in _x("notes/hiring-plan.docx").text
    assert "Revenue grew 12 percent" in _x("board-update.pdf").text


def test_pptx_and_xlsx_roundtrip(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Launch plan"
    prs.save(tmp_path / "d.pptx")
    wb = Workbook()
    wb.active.append(["island", "days"])
    wb.active.append(["Maui", 4])
    wb.save(tmp_path / "t.xlsx")
    assert "Launch plan" in extract_bytes("d.pptx", (tmp_path / "d.pptx").read_bytes()).text
    assert "Maui, 4" in extract_bytes("t.xlsx", (tmp_path / "t.xlsx").read_bytes()).text


def test_truncation_and_unsupported():
    ex = extract_bytes("a.txt", b"x" * 5000, max_chars=1000)
    assert ex.truncated and len(ex.text) == 1000
    assert not is_supported("diagram.png")
    assert extract_bytes("diagram.png", b"\x89PNG").supported is False


def test_corrupt_pdf_is_metadata_only():
    ex = extract_bytes("bad.pdf", b"not a pdf")
    assert ex.supported is False and ex.error
