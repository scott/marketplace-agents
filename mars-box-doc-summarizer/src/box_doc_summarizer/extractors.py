"""Text extraction per file type, with char caps and metadata-only fallback."""

from __future__ import annotations

import io
import os
from dataclasses import dataclass

TEXT_EXTS = {".md", ".markdown", ".txt", ".csv", ".json", ".yaml", ".yml", ".tsv", ".log"}
SUPPORTED_EXTS = TEXT_EXTS | {".pdf", ".docx", ".pptx", ".xlsx"}


def max_chars_per_doc() -> int:
    try:
        return max(500, int(os.environ.get("MAX_CHARS_PER_DOC", "20000")))
    except ValueError:
        return 20000


@dataclass
class Extracted:
    text: str
    supported: bool
    truncated: bool = False
    error: str = ""


def _ext(name: str) -> str:
    return os.path.splitext(name.lower())[1]


def is_supported(name: str) -> bool:
    return _ext(name) in SUPPORTED_EXTS


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"Slide {i}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"Sheet {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(", ".join(cells))
    return "\n".join(parts)


def extract_bytes(name: str, data: bytes, *, max_chars: int | None = None) -> Extracted:
    cap = max_chars or max_chars_per_doc()
    ext = _ext(name)
    if ext not in SUPPORTED_EXTS:
        return Extracted(text="", supported=False)
    try:
        if ext in TEXT_EXTS:
            text = data.decode("utf-8", errors="replace")
        elif ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext == ".pptx":
            text = _pptx(data)
        else:
            text = _xlsx(data)
    except Exception as exc:  # noqa: BLE001
        return Extracted(text="", supported=False, error=f"{type(exc).__name__}: {exc}"[:200])
    text = text.replace("\x00", "").strip()
    truncated = len(text) > cap
    return Extracted(text=text[:cap], supported=True, truncated=truncated)
